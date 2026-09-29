#!/usr/bin/env python3
"""Import NFF XLSX exports into the canonical fixture schedule.

Usage: python scripts/import_senior_xlsx.py /path/to/xlsx-directory
The inputs are not committed; the normalized schedule is.
"""
import csv
import gzip
import json
import re
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SCHEDULE = ROOT / "data/fixture_schedule.csv.gz"
CONFIG = ROOT / "config/competitions.json"

# File prefix, key, display name, NFF tournament ID, gender, district, local teams.
SOURCES = [
    ("2 divisjon menn", "2-div-menn-avd-1", "2. divisjon menn avd. 1", "206007", "MENN", "Nasjonalt", {"Brattvåg"}),
    ("4. div. Menn", "4-div-menn-sunnmore", "4. div. menn", "205682", "MENN", "Sunnmøre", {"Stordal", "Ravn", "Brattvåg 2"}),
    ("4. div kvinner", "4-div-kvinner-sunnmore", "4. div. kvinner", "205920", "KVINNER", "Sunnmøre", {"HaNo", "Brattvåg/Norborg/Ravn"}),
    ("5. div kvinner vår", "5-div-kvinner-var", "5. div. kvinner vår", "207721", "KVINNER", "Sunnmøre", {"Stordal/Ørskog/Skodje"}),
    ("5. div kvinner høst A", "5-div-kvinner-host-a", "5. div. kvinner høst A", "211227", "KVINNER", "Sunnmøre", {"Stordal/Ørskog/Skodje"}),
    ("5. div Menn", "5-div-menn-sunnmore", "5. div. menn", "205921", "MENN", "Sunnmøre", {"HaNo"}),
    ("6. div. menn avd. 2 vår", "6-div-menn-var-2", "6. div. menn vår avd. 2", "207728", "MENN", "Sunnmøre", {"Ørskog/Stordal 2", "Valldal"}),
    ("6. div. menn avd. 3 vår", "6-div-menn-var-3", "6. div. menn vår avd. 3", "207729", "MENN", "Sunnmøre", {"Norborg", "Ravn 2", "Harøy/Lepsøy/HaNo 2", "Skodje"}),
    ("6. div. menn høst", "6-div-menn-host", "6. div. menn høst", "209877", "MENN", "Sunnmøre", {"Norborg"}),
    ("7. div. høst", "7-div-menn-host-1", "7. div. menn høst avd. 1", "210497", "MENN", "Sunnmøre", {"Harøy/Lepsøy/HaNo 2", "Valldal", "Ørskog/Stordal 2", "Ravn 2", "Brattvåg 3", "Skodje"}),
    ("5.div Menn", "5-div-menn-romsdal", "5. div. menn", "205959", "MENN", "Nordmøre og Romsdal", {"Vestnes Varfjell", "Måndalen"}),
    ("6.div Menn avd 03", "6-div-menn-romsdal-3", "6. div. menn avd. 3", "207814", "MENN", "Nordmøre og Romsdal", {"Tomrefjord"}),
    ("Menn 7er", "menn-7er-romsdal", "Menn 7-er", "207913", "MENN", "Nordmøre og Romsdal", {"Fiksdal/Rekdal", "Vågstranda"}),
    ("Kvinner 7er", "kvinner-7er-romsdal", "Kvinner 7-er", "208149", "KVINNER", "Nordmøre og Romsdal", {"Tomrefjord"}),
]


def main(input_dir):
    directory = Path(input_dir)
    with gzip.open(SCHEDULE, "rt", encoding="utf-8", newline="") as fh:
        reader = csv.DictReader(fh)
        fields = reader.fieldnames
        youth = [r for r in reader if r["age"] not in {"MENN", "KVINNER"}]

    config = json.loads(CONFIG.read_text(encoding="utf-8"))
    config["competitions"] = [c for c in config["competitions"] if c["key"] in {r["competitionKey"] for r in youth}]
    senior = []
    seen = {r["matchNumber"] for r in youth}
    for prefix, key, name, tournament, gender, district, local_names in SOURCES:
        candidates = [p for p in directory.glob("*.xlsx") if p.name.startswith(prefix) and ("Romsdal" in p.name) == (district == "Nordmøre og Romsdal")]
        if key == "4-div-menn-sunnmore" and len(candidates) > 1:
            # NFF may also export one club's fixtures under the same series name.
            candidates = [max(candidates, key=lambda p: load_workbook(p, read_only=True).active.max_row)]
        if len(candidates) != 1:
            raise ValueError(f"Expected exactly one XLSX for {key}: {candidates}")
        worksheet = load_workbook(candidates[0], read_only=True, data_only=True).active
        headers, *records = worksheet.values
        if tuple(headers[:11]) != ("Runde", "Dato", "Dag", "Tid", "Hjemmelag", "Resultat", "Bortelag", "Bane", "Turnering", "Kampnummer", "Spillform"):
            raise ValueError(f"Unexpected columns in {candidates[0]}")
        count = 0
        for round_, day, _, time, home, score, away, venue, competition, number, form in records:
            # The Stordal export also contains cup matches. Only league fixtures belong here.
            if key == "4-div-menn-sunnmore" and str(competition).strip() != "4. div. Menn":
                continue
            number = str(number or "").strip()
            if not number.isdigit() or number in seen:
                raise ValueError(f"Missing/duplicate match number: {number} ({key})")
            seen.add(number)
            home, away = str(home).strip(), str(away).strip()
            sides = [side for side, team in (("home", home), ("away", away)) if team in local_names]
            result = re.fullmatch(r"\s*(\d+)\s*[-–]\s*(\d+)\s*", str(score or ""))
            senior.append({
                "competitionKey": f"senior-2026-{key}", "competition": name, "age": gender,
                "round": str(round_ or ""), "date": day.date().isoformat(), "time": str(time or day.strftime("%H:%M")),
                "home": home, "away": away, "venue": str(venue or "").strip(), "format": str(form or ""),
                "matchNumber": number, "seedResult": f"{result[1]}-{result[2]}" if result else "",
                "detailLevel": "full" if sides else "result_only",
                "localTeam": "both" if len(sides) == 2 else (sides[0] if sides else ""),
                "tournamentId": tournament, "discoveryMatchFiksId": "",
            })
            count += 1
        if not count:
            raise ValueError(f"No league fixtures found in {candidates[0]}")
        config["competitions"].append({
            "key": f"senior-2026-{key}", "name": name, "age": gender,
            "season": 2026, "district": district, "tournamentId": tournament,
            "discoveryMatchFiksId": None,
        })
        print(f"{key}: {count} fixtures, {sum(bool(r['localTeam']) for r in senior if r['competitionKey'] == f'senior-2026-{key}')} local")

    with gzip.open(SCHEDULE, "wt", encoding="utf-8", newline="") as fh:
        writer = csv.DictWriter(fh, fieldnames=fields)
        writer.writeheader()
        writer.writerows(youth + senior)
    CONFIG.write_text(json.dumps(config, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Imported {len(senior)} senior fixtures")


if __name__ == "__main__":
    main(sys.argv[1])
